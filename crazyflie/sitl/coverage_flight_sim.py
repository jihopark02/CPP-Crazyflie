#!/usr/bin/env python3
"""
Gazebo 시뮬레이션 -- 경로 오차 개선판 유도 (코너 감속 제외)
ROS2/시각화 하네스와 공통 path_guidance 모듈의 PathGuidance를 결합한다.
  1. 가장 가까운 경로 '점' 대신 경로 '선분'에 현재 위치를 투영
  2. 경로 접선 속도 + 횡오차 보정 속도를 동시에 명령 (linear.y로 직접 횡이동)
  3. 코너 자동감속은 포함하지 않음 -- 순항속도 V 항상 고정
"""
import csv
import math
import os
import sys
import time
import rclpy
from rclpy.node import Node
from nav_msgs.msg import Odometry
from geometry_msgs.msg import Twist
import matplotlib
matplotlib.use('TkAgg')
import matplotlib.pyplot as plt

CRAZYFLIE_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if CRAZYFLIE_ROOT not in sys.path:
    sys.path.insert(0, CRAZYFLIE_ROOT)

from path_guidance import PathGuidance  # noqa: E402

ROBOT_NS = 'crazyflie'
PATH_CSV = sys.argv[1] if len(sys.argv) > 1 else os.path.expanduser(
    '~/CPP-Crazyflie/crazyflie/real-flight/paths/path_obs_1.xlsx')
LOG_CSV = sys.argv[2] if len(sys.argv) > 2 else os.path.expanduser(
    '~/CPP-Crazyflie/crazyflie/sitl/results/obs1/flight_log_obs1.csv')
PLOT_PNG = sys.argv[3] if len(sys.argv) > 3 else None
# 시각화용 장애물 표시 (x,y,size) -- 실제 충돌은 Gazebo world의 obstacle_* 모델이 담당, 이건 그림용
OBSTACLES = [tuple(float(v) for v in o.split(',')) for o in sys.argv[4:]] if len(sys.argv) > 4 else []
ORIGIN_X = 0.0
ORIGIN_Y = 0.0
V = 0.2
L1 = 0.15
GOAL_RADIUS = 0.15
CONTROL_HZ = 20.0
HOVER_HEIGHT = 0.5
FORWARD_SEARCH_DIST_M = 0.35

# 경로 오차 보정 튜닝값 (direct_cflib_coverage_refined2.py와 동일)
CROSS_TRACK_KP = 1.3       # [1/s]
MAX_CROSS_SPEED = 0.10     # [m/s]
MAX_YAW_RATE_DEG = 90.0    # [deg/s]
YAW_KP = 1.8               # [1/s]
DEBUG = False


def clamp(value, low, high):
    return max(low, min(high, value))


def _pick(header, *names):
    for name in names:
        if name in header:
            return header.index(name)
    return None


def load_path(path_file, origin_x, origin_y):
    pts, dists = [], []
    if path_file.lower().endswith('.xlsx'):
        import openpyxl
        wb = openpyxl.load_workbook(path_file, data_only=True)
        ws = wb['path'] if 'path' in wb.sheetnames else wb.worksheets[0]
        rows = list(ws.iter_rows(values_only=True))
        header = list(rows[0])
        ix = _pick(header, 'x', 'x_m')
        iy = _pick(header, 'y', 'y_m')
        idist = _pick(header, 'dist', 'dist_along_path_m')
        for row in rows[1:]:
            pts.append((origin_x + float(row[ix]), origin_y + float(row[iy])))
            dists.append(float(row[idist]) if idist is not None else None)
    else:
        with open(path_file) as f:
            for row in csv.DictReader(f):
                pts.append((origin_x + float(row['x']), origin_y + float(row['y'])))
                dists.append(float(row['dist']) if row.get('dist') not in (None, '') else None)
    if any(d is None for d in dists):
        dists = [0.0]
        for i in range(1, len(pts)):
            dists.append(dists[-1] + math.hypot(pts[i][0]-pts[i-1][0], pts[i][1]-pts[i-1][1]))
    return pts, dists


def wrap_to_pi(a):
    return (a + math.pi) % (2 * math.pi) - math.pi


def quat_to_yaw(q):
    siny_cosp = 2 * (q.w * q.z + q.x * q.y)
    cosy_cosp = 1 - 2 * (q.y * q.y + q.z * q.z)
    return math.atan2(siny_cosp, cosy_cosp)


class CoverageGuidanceSim(Node):
    def __init__(self):
        super().__init__('coverage_guidance_sim_refined')
        self.path, self.dist = load_path(PATH_CSV, ORIGIN_X, ORIGIN_Y)
        self.n = len(self.path)
        self.guidance = PathGuidance(
            self.path,
            self.dist,
            lookahead_distance=L1,
            forward_search_distance=FORWARD_SEARCH_DIST_M,
        )
        self.pos = None
        self.z = 0.0
        self.yaw = 0.0
        self.airborne = False
        self.done = False
        self.t0 = time.time()
        self.last_target = None
        self.last_projection = None
        self.trail_x = []
        self.trail_y = []

        self.get_logger().info(
            f'경로 로드: {self.n}개 점, 총 길이 {self.dist[-1]:.2f}m, 원점({ORIGIN_X:.2f},{ORIGIN_Y:.2f})')

        self.log_file = open(LOG_CSV, 'w', newline='')
        self.log_writer = csv.writer(self.log_file)
        self.log_writer.writerow(['t', 'x', 'y', 'z', 'yaw_deg', 'path_idx',
                                   'target_x', 'target_y', 'eta_deg', 'cross_track_err'])

        self.create_subscription(Odometry, f'/{ROBOT_NS}/odom', self.odom_cb, 10)
        self.pub = self.create_publisher(Twist, '/cmd_vel', 10)

    def odom_cb(self, msg):
        self.pos = (msg.pose.pose.position.x, msg.pose.pose.position.y)
        self.z = msg.pose.pose.position.z
        self.yaw = quat_to_yaw(msg.pose.pose.orientation)
        if self.airborne:
            self.trail_x.append(self.pos[0])
            self.trail_y.append(self.pos[1])

    def finish(self, reason):
        self.done = True
        self.publish_vel(0.0, 0.0, 0.0)
        self.log_file.close()
        self.get_logger().info(f'{reason} -- 로그 저장됨: {LOG_CSV}')

    def control_loop(self):
        if self.pos is None or self.done:
            return

        if not self.airborne:
            if self.z < HOVER_HEIGHT:
                self.publish_vel(0.0, 0.0, 0.0, vz=0.5)
                return
            else:
                self.airborne = True

        x, y = self.pos
        gx, gy = self.path[-1]
        if (self.guidance.path_idx >= self.n - 5 and
                math.hypot(x - gx, y - gy) < GOAL_RADIUS):
            self.finish('커버리지 경로 완주')
            return

        proj_x, proj_y, ref_x, ref_y, tx, ty, _ = self.guidance.update(x, y)
        self.last_target = (ref_x, ref_y)
        self.last_projection = (proj_x, proj_y)

        # 횡오차 P 보정 (코너 감속 없이 순항속도 V는 항상 고정)
        corr_x = CROSS_TRACK_KP * (proj_x - x)
        corr_y = CROSS_TRACK_KP * (proj_y - y)
        corr_mag = math.hypot(corr_x, corr_y)
        if corr_mag > MAX_CROSS_SPEED:
            scale = MAX_CROSS_SPEED / corr_mag
            corr_x *= scale
            corr_y *= scale

        vx_world = V * tx + corr_x
        vy_world = V * ty + corr_y

        planar_speed = math.hypot(vx_world, vy_world)
        max_planar_speed = math.hypot(V, MAX_CROSS_SPEED)
        if planar_speed > max_planar_speed:
            scale = max_planar_speed / planar_speed
            vx_world *= scale
            vy_world *= scale

        cos_yaw, sin_yaw = math.cos(self.yaw), math.sin(self.yaw)
        vx_body = cos_yaw * vx_world + sin_yaw * vy_world
        vy_body = -sin_yaw * vx_world + cos_yaw * vy_world

        desired_yaw = math.atan2(ty, tx)
        yaw_error = wrap_to_pi(desired_yaw - self.yaw)
        yaw_rate = clamp(YAW_KP * yaw_error,
                          -math.radians(MAX_YAW_RATE_DEG), math.radians(MAX_YAW_RATE_DEG))

        cross_track_err = math.hypot(proj_x - x, proj_y - y)
        eta = wrap_to_pi(math.atan2(ref_y - y, ref_x - x) - self.yaw)

        t = time.time() - self.t0
        self.log_writer.writerow([f'{t:.2f}', f'{x:.4f}', f'{y:.4f}',
                                   f'{self.z:.4f}', f'{math.degrees(self.yaw):.2f}',
                                   self.guidance.path_idx, f'{ref_x:.4f}', f'{ref_y:.4f}',
                                   f'{math.degrees(eta):.2f}', f'{cross_track_err:.4f}'])

        if DEBUG:
            self.get_logger().info(
                f'pos=({x:.2f},{y:.2f}) idx={self.guidance.path_idx}/{self.n} '
                f'e={cross_track_err:.3f} vx={vx_body:.2f} vy={vy_body:.2f}')

        self.publish_vel(vx_body, vy_body, yaw_rate)

    def publish_vel(self, vx, vy, wz, vz=0.0):
        msg = Twist()
        msg.linear.x = vx
        msg.linear.y = vy
        msg.linear.z = vz
        msg.angular.z = wz
        self.pub.publish(msg)


def main():
    rclpy.init()
    node = CoverageGuidanceSim()

    plt.ion()
    fig, ax = plt.subplots(figsize=(8, 7))
    path_x = [p[0] for p in node.path]
    path_y = [p[1] for p in node.path]
    ax.plot(path_x, path_y, 'b-', linewidth=1.0, alpha=0.5, label='planned path')
    trail_line, = ax.plot([], [], '-', color='gray', linewidth=1.5, label='flown trail')
    drone_dot, = ax.plot([], [], 'ko', markersize=10, label='drone')
    target_dot, = ax.plot([], [], 'ro', markersize=8, label='lookahead target')
    projection_dot, = ax.plot([], [], 'go', markersize=5, label='path projection')
    for i, (ox, oy, osize) in enumerate(OBSTACLES):
        ax.add_patch(plt.Rectangle((ox - osize / 2, oy - osize / 2), osize, osize,
                                    facecolor='red', alpha=0.4, edgecolor='darkred',
                                    label='obstacle' if i == 0 else None))
    ax.set_xlim(min(path_x) - 0.3, max(path_x) + 0.3)
    ax.set_ylim(min(path_y) - 0.3, max(path_y) + 0.3)
    ax.set_aspect('equal')
    ax.set_xlabel('x (m)')
    ax.set_ylabel('y (m)')
    ax.legend(loc='upper right')
    ax.set_title('Coverage path -- cross-track feedback (sim)')

    last_control_time = time.time()
    control_period = 1.0 / CONTROL_HZ

    try:
        while rclpy.ok():
            rclpy.spin_once(node, timeout_sec=0.01)

            now = time.time()
            if now - last_control_time >= control_period:
                node.control_loop()
                last_control_time = now

            if node.pos is not None:
                drone_dot.set_data([node.pos[0]], [node.pos[1]])
            if node.last_target is not None:
                target_dot.set_data([node.last_target[0]], [node.last_target[1]])
            if node.last_projection is not None:
                projection_dot.set_data([node.last_projection[0]], [node.last_projection[1]])
            if node.trail_x:
                trail_line.set_data(node.trail_x, node.trail_y)

            fig.canvas.draw_idle()
            fig.canvas.flush_events()
            plt.pause(0.001)

            if node.done:
                if PLOT_PNG:
                    fig.savefig(PLOT_PNG, dpi=150)
                    print(f'궤적 그래프 저장됨: {PLOT_PNG}')
                break
    except KeyboardInterrupt:
        if not node.log_file.closed:
            node.log_file.close()
    finally:
        node.destroy_node()
        rclpy.shutdown()
        plt.ioff()
        plt.show()


if __name__ == '__main__':
    main()
