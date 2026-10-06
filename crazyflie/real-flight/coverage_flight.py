#!/usr/bin/env python3
"""
cf231 실비행 -- 경로 오차 개선판

기존 direct_cflib_coverage.py의 통신/추정기/로그/이착륙 구조는 유지한다.
유도부만 다음과 같이 변경했다.
  1. 가장 가까운 경로 '점' 대신 경로 '선분'에 현재 위치를 투영
  2. 경로 접선 속도 + 횡오차 보정 속도를 동시에 명령
  3. world-frame 속도를 hover setpoint의 body-frame vx/vy로 변환
  4. 전방 경로의 방향 변화가 크면 자동 감속
"""
import sys
import os
import csv
import math
import time
import threading
import matplotlib

matplotlib.use('TkAgg')
import matplotlib.pyplot as plt  # noqa: E402
from matplotlib.ticker import MultipleLocator  # noqa: E402

CRAZYFLIE_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if CRAZYFLIE_ROOT not in sys.path:
    sys.path.insert(0, CRAZYFLIE_ROOT)

from path_guidance import PathGuidance  # noqa: E402

sys.path.insert(0, os.path.expanduser(
    '~/natnet_ws/src/natnet_ros2/deps/NatNetSDK/samples/PythonClient'))
from NatNetClient import NatNetClient  # noqa: E402

import cflib.crtp
from cflib.crazyflie import Crazyflie
from cflib.crazyflie.syncCrazyflie import SyncCrazyflie
from cflib.crazyflie.syncLogger import SyncLogger
from cflib.crazyflie.log import LogConfig

URI = 'radio://0/80/2M/E7E7E7E7E7'
SERVER_IP = '192.168.50.49'
CLIENT_IP = '192.168.50.236'
USE_MULTICAST = False

PATH_CSV = os.path.expanduser('~/natnet_ws/coverage_path_obs1_50cm.csv')
LOG_CSV = os.path.expanduser('~/natnet_ws/flight_log.csv')
OUT_PNG = os.path.expanduser('~/natnet_ws/flight_error_plot.png')
V = 0.2
L1 = 0.15
GOAL_RADIUS = 0.15
FORWARD_SEARCH_DIST_M = 0.35
HOVER_HEIGHT = 0.5
CONTROL_HZ = 20.0
PLOT_HZ = 5.0
PRINT_HZ = 2.0

# 경로 오차 보정 튜닝값. 처음에는 이 세 값만 조정하면 된다.
CROSS_TRACK_KP = 1.3       # [1/s], 클수록 경로로 강하게 복귀
MAX_CROSS_SPEED = 0.10     # [m/s], 횡방향 보정 속도 제한
MIN_TURN_SPEED = 0.05      # [m/s], U턴 구간 최소 전진 속도
MAX_YAW_RATE_DEG = 90.0    # [deg/s]
YAW_KP = 1.8               # [1/s]

# 경로 진입 단계: 첫 경로점 근처로 이동한 뒤 본 추종을 시작한다.
APPROACH_SPEED = 0.08       # [m/s]
START_TOLERANCE = 0.03      # [m]
APPROACH_TIMEOUT = 12.0     # [s]

# 이 정도의 선행 방향 변화에서 U턴 최소 속도까지 감속한다.
TURN_FULL_SLOWDOWN_DEG = 20.0

MAX_TILT_DEG = 25.0
MIN_VBAT = 3.7
MAX_Z = 1.2


def clamp(value, low, high):
    return max(low, min(high, value))


def _pick(header, *names):
    for name in names:
        if name in header:
            return header.index(name)
    return None


def load_path(path_file):
    pts, dists = [], []
    if path_file.lower().endswith('.xlsx'):
        import openpyxl
        wb = openpyxl.load_workbook(path_file, data_only=True)
        ws = wb['path'] if 'path' in wb.sheetnames else wb.worksheets[0]
        rows = list(ws.iter_rows(values_only=True))
        header = list(rows[0])
        ix = _pick(header, 'x', 'x_m')
        iy = _pick(header, 'y', 'y_m')
        idist = _pick(header, 'dist', 'dist_along_path_m')
        for row in rows[1:]:
            pts.append((float(row[ix]), float(row[iy])))
            dists.append(float(row[idist]) if idist is not None else None)
    else:
        with open(path_file) as f:
            for row in csv.DictReader(f):
                pts.append((float(row['x']), float(row['y'])))
                dists.append(float(row['dist']))
    if any(d is None for d in dists):
        dists = [0.0]
        for i in range(1, len(pts)):
            dists.append(dists[-1] + math.hypot(pts[i][0]-pts[i-1][0], pts[i][1]-pts[i-1][1]))
    if len(pts) < 2:
        raise ValueError('경로에는 최소 2개 점이 필요합니다')
    return pts, dists


def wrap_to_pi(a):
    return (a + math.pi) % (2 * math.pi) - math.pi


class MocapState:
    def __init__(self):
        self.latest = None
        self.first_id = None
        self.lock = threading.Lock()

    def rigid_body_cb(self, new_id, pos, rot):
        with self.lock:
            if self.first_id is None:
                self.first_id = new_id
                print(f'[mocap] 첫 Rigid Body ID={new_id} 를 드론으로 사용')
            if new_id == self.first_id:
                self.latest = (pos[0], pos[1], pos[2], rot[0], rot[1], rot[2], rot[3])


def start_natnet(mocap_state):
    client = NatNetClient()
    client.set_client_address(CLIENT_IP)
    client.set_server_address(SERVER_IP)
    client.set_use_multicast(USE_MULTICAST)
    client.rigid_body_listener = mocap_state.rigid_body_cb
    ok = client.run('d')
    if not ok:
        raise RuntimeError('NatNet 클라이언트 시작 실패')
    return client


def setup_estimator_params(cf):
    cf.param.set_value('stabilizer.estimator', '2')
    time.sleep(0.1)
    cf.param.set_value('locSrv.extPosStdDev', '1e-3')
    time.sleep(0.1)
    cf.param.set_value('locSrv.extQuatStdDev', '0.5e-1')
    time.sleep(0.1)
    print('추정기 파라미터 설정 완료 (estimator=kalman, extPosStdDev=1e-3, extQuatStdDev=0.5e-1)')


def reset_estimator(cf):
    cf.param.set_value('kalman.resetEstimation', '1')
    time.sleep(0.1)
    cf.param.set_value('kalman.resetEstimation', '0')
    wait_for_position_estimator(cf)


def wait_for_position_estimator(cf, timeout=10.0):
    log_config = LogConfig(name='KalmanVar', period_in_ms=200)
    log_config.add_variable('kalman.varPX', 'float')
    log_config.add_variable('kalman.varPY', 'float')
    log_config.add_variable('kalman.varPZ', 'float')

    var_x_hist = [1000] * 10
    var_y_hist = [1000] * 10
    var_z_hist = [1000] * 10
    threshold = 0.001

    t0 = time.time()
    with SyncLogger(cf, log_config) as logger:
        for entry in logger:
            data = entry[1]
            var_x_hist.append(data['kalman.varPX']); var_x_hist.pop(0)
            var_y_hist.append(data['kalman.varPY']); var_y_hist.pop(0)
            var_z_hist.append(data['kalman.varPZ']); var_z_hist.pop(0)

            if (max(var_x_hist) - min(var_x_hist) < threshold and
                    max(var_y_hist) - min(var_y_hist) < threshold and
                    max(var_z_hist) - min(var_z_hist) < threshold):
                print('추정기 수렴 완료')
                return True
            if time.time() - t0 > timeout:
                print('수렴 타임아웃 -- 그래도 진행')
                return False




def main():
    mocap = MocapState()
    print('NatNet 연결 시도...')
    natnet_client = start_natnet(mocap)

    print('mocap 데이터 대기 중...')
    t_wait0 = time.time()
    while mocap.latest is None:
        time.sleep(0.05)
        if time.time() - t_wait0 > 10.0:
            print('mocap 데이터가 10초 넘게 안 들어옴 -- Motive에서 cf231 Solved 상태인지 확인하세요')
            return
    print('mocap 데이터 수신 시작:', mocap.latest)

    path, dist = load_path(PATH_CSV)
    print(f'경로 로드: {len(path)}개 점, 총 길이 {dist[-1]:.2f}m')
    guidance = PathGuidance(
        path,
        dist,
        lookahead_distance=L1,
        forward_search_distance=FORWARD_SEARCH_DIST_M,
    )

    log_file = open(LOG_CSV, 'w', newline='')
    log_writer = csv.writer(log_file)
    log_writer.writerow(['t', 'x', 'y', 'z', 'yaw_deg', 'path_idx',
                          'target_x', 'target_y', 'eta_deg', 'cross_track_err',
                          'signed_error'])

    plt.ion()
    fig, ax = plt.subplots(figsize=(8, 7))
    path_x = [p[0] for p in path]
    path_y = [p[1] for p in path]
    ax.plot(path_x, path_y, 'b-', linewidth=1.0, alpha=0.5, label='planned path')
    trail_line, = ax.plot([], [], '-', color='gray', linewidth=1.5, label='flown trail')
    drone_dot, = ax.plot([], [], 'ko', markersize=10, label='drone')
    target_dot, = ax.plot([], [], 'ro', markersize=8, label='lookahead target')
    projection_dot, = ax.plot([], [], 'go', markersize=5, label='path projection')
    margin = 0.3
    ax.set_xlim(min(path_x) - margin, max(path_x) + margin)
    ax.set_ylim(min(path_y) - margin, max(path_y) + margin)
    ax.set_aspect('equal')
    ax.set_xlabel('x (m)')
    ax.set_ylabel('y (m)')
    ax.legend(loc='upper right')
    ax.set_title('cf231 coverage path -- cross-track feedback')
    trail_x, trail_y = [], []

    def update_plot():
        trail_x.append(state['x'])
        trail_y.append(state['y'])
        drone_dot.set_data([state['x']], [state['y']])
        trail_line.set_data(trail_x, trail_y)
        fig.canvas.draw_idle()
        fig.canvas.flush_events()

    cflib.crtp.init_drivers()

    with SyncCrazyflie(URI, cf=Crazyflie(rw_cache='./cache')) as scf:
        cf = scf.cf
        setup_estimator_params(cf)

        print('아밍...')
        cf.platform.send_arming_request(True)
        time.sleep(1.0)

        stop_flag = threading.Event()
        fault = threading.Event()

        def extpose_loop():
            while not stop_flag.is_set():
                if mocap.latest is not None:
                    x, y, z, qx, qy, qz, qw = mocap.latest
                    cf.extpos.send_extpose(x, y, z, qx, qy, qz, qw)
                time.sleep(0.02)

        extpose_thread = threading.Thread(target=extpose_loop, daemon=True)
        extpose_thread.start()

        time.sleep(1.0)
        print('추정기 리셋 + 수렴 대기...')
        reset_estimator(cf)

        state = {'x': 0.0, 'y': 0.0, 'z': 0.0, 'yaw': 0.0,
                 'roll': 0.0, 'pitch': 0.0, 'vbat': 4.2}

        def pose_log_cb(timestamp, data, logconf):
            state['x'] = data['stateEstimate.x']
            state['y'] = data['stateEstimate.y']
            state['z'] = data['stateEstimate.z']
            state['yaw'] = data['stabilizer.yaw']

        def safety_log_cb(timestamp, data, logconf):
            state['roll'] = data['stabilizer.roll']
            state['pitch'] = data['stabilizer.pitch']
            state['vbat'] = data['pm.vbat']
            if not fault.is_set() and (
                    abs(state['roll']) > MAX_TILT_DEG or
                    abs(state['pitch']) > MAX_TILT_DEG or
                    state['vbat'] < MIN_VBAT or
                    state['z'] > MAX_Z):
                fault.set()
                print(f"비상 감지: roll={state['roll']:.1f} pitch={state['pitch']:.1f} "
                      f"vbat={state['vbat']:.2f} z={state['z']:.2f}")

        pose_log = LogConfig(name='PoseMonitor', period_in_ms=50)
        pose_log.add_variable('stateEstimate.x', 'float')
        pose_log.add_variable('stateEstimate.y', 'float')
        pose_log.add_variable('stateEstimate.z', 'float')
        pose_log.add_variable('stabilizer.yaw', 'float')
        cf.log.add_config(pose_log)
        pose_log.data_received_cb.add_callback(pose_log_cb)
        pose_log.start()

        safety_log = LogConfig(name='SafetyMonitor', period_in_ms=50)
        safety_log.add_variable('stabilizer.roll', 'float')
        safety_log.add_variable('stabilizer.pitch', 'float')
        safety_log.add_variable('pm.vbat', 'float')
        cf.log.add_config(safety_log)
        safety_log.data_received_cb.add_callback(safety_log_cb)
        safety_log.start()

        def emergency_stop(reason):
            print(f'비상 정지: {reason}')
            cf.commander.send_velocity_world_setpoint(0, 0, 0, 0)
            time.sleep(0.1)
            cf.platform.send_arming_request(False)

        dt = 1.0 / CONTROL_HZ

        try:
            print('이륙...')
            t_takeoff0 = time.time()
            target_z = 0.0
            while True:
                if fault.is_set():
                    emergency_stop('이륙 중 이상 감지')
                    return
                if time.time() - t_takeoff0 > 8.0:
                    emergency_stop('이륙 타임아웃')
                    return
                target_z = min(HOVER_HEIGHT, target_z + 0.3 * dt)
                cf.commander.send_hover_setpoint(0, 0, 0, target_z)
                if target_z >= HOVER_HEIGHT and state['z'] > HOVER_HEIGHT - 0.05:
                    break
                update_plot()
                plt.pause(dt)

            # 본 추종 오차에 진입 구간이 섞이지 않도록 첫 경로점으로 정렬한다.
            print('경로 시작점 정렬...')
            start_x, start_y = path[0]
            first_dx = path[1][0] - start_x
            first_dy = path[1][1] - start_y
            first_heading = math.atan2(first_dy, first_dx)
            t_approach0 = time.time()
            while True:
                if fault.is_set():
                    emergency_stop('경로 진입 중 이상 감지')
                    return
                if time.time() - t_approach0 > APPROACH_TIMEOUT:
                    emergency_stop('경로 시작점 진입 타임아웃')
                    return

                x, y = state['x'], state['y']
                yaw = math.radians(state['yaw'])
                dx, dy = start_x - x, start_y - y
                distance_to_start = math.hypot(dx, dy)
                if distance_to_start <= START_TOLERANCE:
                    break

                approach_v = min(APPROACH_SPEED, max(0.025, distance_to_start))
                vx_world = approach_v * dx / distance_to_start
                vy_world = approach_v * dy / distance_to_start
                cos_yaw = math.cos(yaw)
                sin_yaw = math.sin(yaw)
                vx_body = cos_yaw * vx_world + sin_yaw * vy_world
                vy_body = -sin_yaw * vx_world + cos_yaw * vy_world
                yaw_error = wrap_to_pi(first_heading - yaw)
                yaw_rate = clamp(
                    math.degrees(YAW_KP * yaw_error),
                    -MAX_YAW_RATE_DEG,
                    MAX_YAW_RATE_DEG,
                )
                cf.commander.send_hover_setpoint(
                    vx_body, vy_body, yaw_rate, HOVER_HEIGHT)
                update_plot()
                plt.pause(dt)

            cf.commander.send_hover_setpoint(0, 0, 0, HOVER_HEIGHT)
            time.sleep(0.3)

            # 진입 궤적을 본 경로 추종 궤적/통계에서 제외한다.
            trail_x.clear()
            trail_y.clear()
            trail_line.set_data([], [])
            error_time = []
            signed_error_history = []
            altitude_error_history = []

            print('선분 투영 + 횡오차 피드백 유도 시작')
            t_mission0 = time.time()
            gx, gy = path[-1]
            next_tick = time.monotonic()
            last_plot_time = next_tick
            last_print_time = next_tick - 1.0 / PRINT_HZ
            while True:
                if fault.is_set():
                    emergency_stop('유도 중 이상 감지')
                    return
                if time.time() - t_mission0 > 180.0:
                    emergency_stop('미션 타임아웃')
                    return

                x, y = state['x'], state['y']
                yaw = math.radians(state['yaw'])
                if (guidance.path_idx >= guidance.n - 5 and
                        math.hypot(x - gx, y - gy) < GOAL_RADIUS):
                    print('경로 완주')
                    break

                (proj_x, proj_y, ref_x, ref_y,
                 tx, ty, turn_angle) = guidance.update(x, y)

                # 전방 방향 변화가 클수록 감속한다. 직선에서는 기존 V를 유지한다.
                turn_ratio = clamp(
                    turn_angle / math.radians(TURN_FULL_SLOWDOWN_DEG),
                    0.0,
                    1.0,
                )
                v_along = V - (V - MIN_TURN_SPEED) * turn_ratio

                # 투영점으로 향하는 P 보정. 크기를 제한해 급격한 횡이동을 막는다.
                corr_x = CROSS_TRACK_KP * (proj_x - x)
                corr_y = CROSS_TRACK_KP * (proj_y - y)
                corr_mag = math.hypot(corr_x, corr_y)
                if corr_mag > MAX_CROSS_SPEED:
                    scale = MAX_CROSS_SPEED / corr_mag
                    corr_x *= scale
                    corr_y *= scale

                vx_world = v_along * tx + corr_x
                vy_world = v_along * ty + corr_y

                # 과도한 합성 속도를 막되, 보정 방향은 유지한다.
                planar_speed = math.hypot(vx_world, vy_world)
                max_planar_speed = math.hypot(V, MAX_CROSS_SPEED)
                if planar_speed > max_planar_speed:
                    scale = max_planar_speed / planar_speed
                    vx_world *= scale
                    vy_world *= scale

                # hover setpoint의 vx/vy는 body frame이므로 world -> body 변환.
                cos_yaw = math.cos(yaw)
                sin_yaw = math.sin(yaw)
                vx_body = cos_yaw * vx_world + sin_yaw * vy_world
                vy_body = -sin_yaw * vx_world + cos_yaw * vy_world

                desired_yaw = math.atan2(ty, tx)
                yaw_error = wrap_to_pi(desired_yaw - yaw)
                yaw_rate = clamp(
                    math.degrees(YAW_KP * yaw_error),
                    -MAX_YAW_RATE_DEG,
                    MAX_YAW_RATE_DEG,
                )

                cf.commander.send_hover_setpoint(
                    vx_body, vy_body, yaw_rate, HOVER_HEIGHT)

                signed_error = tx * (y - proj_y) - ty * (x - proj_x)
                t_now = time.time() - t_mission0
                error_time.append(t_now)
                signed_error_history.append(signed_error)
                altitude_error_history.append(state['z'] - HOVER_HEIGHT)
                log_writer.writerow([t_now, x, y, state['z'], math.degrees(yaw),
                                      guidance.path_idx, ref_x, ref_y,
                                      math.degrees(yaw_error), abs(signed_error),
                                      signed_error])
                log_file.flush()

                now = time.monotonic()
                if now - last_print_time >= 1.0 / PRINT_HZ:
                    print(f'  idx={guidance.path_idx}/{guidance.n} '
                          f'x={x:.2f} y={y:.2f} z={state["z"]:.2f} '
                          f'e={signed_error:+.3f} v={v_along:.2f}')
                    last_print_time = now

                if now - last_plot_time >= 1.0 / PLOT_HZ:
                    target_dot.set_data([ref_x], [ref_y])
                    projection_dot.set_data([proj_x], [proj_y])
                    update_plot()
                    plt.pause(0.001)
                    last_plot_time = now

                # 그래프 처리 시간과 무관하게 제어 주기를 20 Hz에 맞춘다.
                next_tick += dt
                sleep_time = next_tick - time.monotonic()
                if sleep_time > 0:
                    time.sleep(sleep_time)
                else:
                    next_tick = time.monotonic()

            print('착륙...')
            t_land0 = time.time()
            target_z = HOVER_HEIGHT
            while target_z > 0.02:
                if fault.is_set():
                    emergency_stop('착륙 중 이상 감지')
                    return
                if time.time() - t_land0 > 10.0:
                    emergency_stop('착륙 타임아웃')
                    return
                target_z = max(0.0, target_z - 0.2 * dt)
                cf.commander.send_hover_setpoint(0, 0, 0, target_z)
                update_plot()
                plt.pause(dt)

            cf.commander.send_velocity_world_setpoint(0, 0, 0, 0)
            time.sleep(0.1)
            cf.platform.send_arming_request(False)
            print('완료 -- 디스암')

        finally:
            stop_flag.set()
            log_file.close()
            print(f'비행 로그 저장됨: {LOG_CSV}')

    natnet_client.shutdown()

    plt.ioff()

    if error_time:
        error_fig, error_ax = plt.subplots(2, 1, figsize=(9, 6), sharex=True)
        error_ax[0].plot(error_time, signed_error_history, 'b-')
        error_ax[0].axhline(0.0, color='gray', linewidth=0.8)
        error_ax[0].set_ylabel('signed error (m)')
        error_ax[0].set_title('signed cross-track error vs time')
        error_ax[0].set_ylim(-0.1, 0.1)
        error_ax[0].yaxis.set_major_locator(MultipleLocator(0.02))
        error_ax[0].grid(True, alpha=0.3)

        error_ax[1].plot(error_time, altitude_error_history, 'r-')
        error_ax[1].axhline(0.0, color='gray', linewidth=0.8)
        error_ax[1].set_xlabel('time (s)')
        error_ax[1].set_ylabel('altitude error (m)')
        error_ax[1].set_title('altitude error (goal 0.5 m)')
        error_ax[1].grid(True, alpha=0.3)
        error_fig.tight_layout()
        error_fig.savefig(OUT_PNG, dpi=150)
        print(f'오차 그래프 저장됨: {OUT_PNG}')

    print('비행 종료 -- 그래프 창을 닫으면 프로그램이 끝납니다')
    plt.show()


if __name__ == '__main__':
    main()
